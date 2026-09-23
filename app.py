from __future__ import annotations

import csv
import io
import os
import re
from collections import defaultdict
from functools import lru_cache
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

import pandas as pd
from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


PROJECT_ROOT = Path(__file__).resolve().parent
VERSION_ORDER = ("H0", "H1", "H2", "H3", "H4", "H5", "H6")
VERSION_LABELS = {
    "H0": "HS92",
    "H1": "HS96",
    "H2": "HS02",
    "H3": "HS07",
    "H4": "HS12",
    "H5": "HS17",
    "H6": "HS22",
}
VERSION_YEARS = {
    "H0": 1992,
    "H1": 1996,
    "H2": 2002,
    "H3": 2007,
    "H4": 2012,
    "H5": 2017,
    "H6": 2022,
}
SOURCE_LINKS = {
    "un_statistics": "https://unstats.un.org/unsd/classifications/Econ",
    "un_hs_descriptions": "https://unstats.un.org/unsd/classifications/Econ/download/In%20Text/HSCodeandDescription.xlsx",
    "wto_hs_tracker": "https://hstracker.wto.org/",
}
OFFICIAL_DESCRIPTION_WORKBOOK = PROJECT_ROOT / "data" / "official_hs_descriptions" / "HSCodeandDescription.xlsx"


def resolve_weights_dir() -> Path:
    """Use the conversion weights bundled in this project."""
    return PROJECT_ROOT / "data" / "conversion_weights"


WEIGHTS_DIR = resolve_weights_dir()


class ConversionError(ValueError):
    pass


def parse_target_years(value: Any) -> list[int]:
    if value is None:
        raise ConversionError("请输入目标年份。")
    if isinstance(value, list):
        raw_values = value
    else:
        raw_values = re.split(r"[\s,，、;；]+", str(value).strip())
    years: list[int] = []
    for raw_value in raw_values:
        text = str(raw_value).strip()
        if not text:
            continue
        try:
            year = int(text)
        except ValueError as exc:
            raise ConversionError("目标年份请输入整数，例如：2005 2010 2015 2020。") from exc
        year_to_version(year)
        if year not in years:
            years.append(year)
    if not years:
        raise ConversionError("请输入目标年份。")
    return years


def normalize_code(value: str) -> str:
    code = str(value).strip().replace("\ufeff", "")
    if code.endswith(".0"):
        code = code[:-2]
    if not code.isdigit() or len(code) > 6:
        raise ConversionError("HS Code 必须是 1-6 位数字。")
    return code.zfill(6)


def parse_hs_codes(value: Any) -> list[str]:
    if value is None:
        raise ConversionError("请输入至少一个 HS Code。")
    if isinstance(value, list):
        raw_values = value
    else:
        raw_values = re.split(r"[\s,，、;；]+", str(value).strip())
    codes: list[str] = []
    for raw_value in raw_values:
        text = str(raw_value).strip().lstrip("'")
        if not text:
            continue
        try:
            code = normalize_code(text)
        except ConversionError as exc:
            raise ConversionError("HS Code 必须是 1-6 位数字，可输入多个代码。") from exc
        if code not in codes:
            codes.append(code)
    if not codes:
        raise ConversionError("请输入至少一个 HS Code。")
    return codes


def _normalize_uploaded_code(value: Any) -> str | None:
    if pd.isna(value):
        return None
    text = str(value).strip().lstrip("'")
    if re.fullmatch(r"\d+\.0", text):
        text = text[:-2]
    if not re.fullmatch(r"\d{5,6}", text):
        return None
    return text.zfill(6)


def extract_uploaded_hs_codes(frame: pd.DataFrame) -> list[str]:
    if "HSCode" not in frame.columns:
        raise ConversionError("上传文件必须包含列名 HSCode。")
    codes = [_normalize_uploaded_code(value) for value in frame["HSCode"].tolist()]
    return list(dict.fromkeys(code for code in codes if code))


def read_uploaded_hs_codes(file_storage) -> list[str]:
    filename = (file_storage.filename or "").lower()
    if not (filename.endswith(".csv") or filename.endswith(".xlsx")):
        raise ConversionError("请上传 CSV 或 XLSX 格式文件。")

    if filename.endswith(".xlsx"):
        sheet_frames = pd.read_excel(file_storage, sheet_name=None, dtype=str)
    else:
        raw = file_storage.read()
        frame = None
        last_error: Exception | None = None
        for encoding in ("utf-8-sig", "utf-16", "gb18030", "cp1252"):
            try:
                text = raw.decode(encoding)
                try:
                    delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
                except csv.Error:
                    delimiter = ","
                frame = pd.read_csv(io.StringIO(text), sep=delimiter, engine="python", dtype=str)
                break
            except (UnicodeDecodeError, csv.Error, ValueError, pd.errors.ParserError) as exc:
                last_error = exc
        if frame is None:
            detail = str(last_error) if last_error else "未知错误"
            raise ConversionError(f"CSV 无法解析：{detail}")
        sheet_frames = {"csv": frame}

    codes: list[str] = []
    matched_sheet = False
    for frame in sheet_frames.values():
        if "HSCode" not in frame.columns:
            continue
        matched_sheet = True
        for code in extract_uploaded_hs_codes(frame):
            if code not in codes:
                codes.append(code)
    if not matched_sheet:
        raise ConversionError("上传文件必须包含列名 HSCode。")
    if not codes:
        raise ConversionError("HSCode 列中未找到有效的六位 HS Code。")
    return codes


def year_to_version(year: int) -> str:
    eligible = [version for version in VERSION_ORDER if VERSION_YEARS[version] <= year]
    if not eligible:
        raise ConversionError("目标年份不能早于 1992。")
    return eligible[-1]


def version_from_value(value: str | None) -> str | None:
    if not value or value.strip().upper() in {"AUTO", "自动", "自动识别"}:
        return None
    candidate = value.strip().upper()
    if candidate in VERSION_ORDER:
        return candidate
    label_to_version = {label: version for version, label in VERSION_LABELS.items()}
    if candidate in label_to_version:
        return label_to_version[candidate]
    raise ConversionError("HS 版本必须是 HS92、HS96、HS02、HS07、HS12、HS17 或 HS22。")


def _version_pair(from_version: str, to_version: str) -> tuple[str, str]:
    from_index = VERSION_ORDER.index(from_version)
    to_index = VERSION_ORDER.index(to_version)
    if from_index == to_index:
        raise ValueError("same version")
    return (from_version, to_version)


@lru_cache(maxsize=32)
def load_direct_mapping(from_version: str, to_version: str) -> dict[str, list[dict[str, Any]]]:
    """Load one adjacent Harvard weight table as a sparse code-to-code map.

    The source tables contain conversion relationships rather than a complete
    catalogue of unchanged codes. The service therefore adds an explicit
    identity fallback when a code is absent from a sparse table. The response
    marks that edge as ``implicit_identity`` so it is auditable in the UI.
    """
    _version_pair(from_version, to_version)
    path = WEIGHTS_DIR / f"conversion_weights_{from_version}_to_{to_version}.csv"
    if not path.exists():
        raise ConversionError(f"缺少转换文件：{path}")

    try:
        frame = pd.read_csv(path, dtype=str, usecols=[from_version, to_version, "weight"])
    except Exception as exc:  # pragma: no cover - surfaced as a user-facing API error
        raise ConversionError(f"读取转换文件失败：{path.name}（{exc}）") from exc

    frame[from_version] = frame[from_version].map(normalize_code)
    frame[to_version] = frame[to_version].map(normalize_code)
    frame["weight"] = pd.to_numeric(frame["weight"], errors="coerce").fillna(0.0)
    frame = frame[frame["weight"] > 0].copy()
    frame = frame.groupby([from_version, to_version], as_index=False)["weight"].sum()

    mapping: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in frame.sort_values([from_version, "weight", to_version], ascending=[True, False, True]).itertuples(index=False):
        mapping[getattr(row, from_version)].append(
            {
                "code": getattr(row, to_version),
                "weight": round(float(row.weight), 8),
                "relation": "direct",
            }
        )
    return dict(mapping)


@lru_cache(maxsize=1)
def build_version_membership_index() -> tuple[dict[str, set[str]], dict[tuple[str, str], set[str]]]:
    """Build a version index from positive-weight relationships.

    Harvard's files are sparse: unchanged codes may be omitted. A code that
    appears on one side of a relationship can therefore be carried across an
    adjacent release only when that code has no explicit non-identity change on
    that release boundary. This lets us infer the earliest defensible version
    while keeping the uncertainty visible to the caller.
    """
    observed: dict[str, set[str]] = {version: set() for version in VERSION_ORDER}
    changed: dict[tuple[str, str], set[str]] = {}

    for index in range(len(VERSION_ORDER) - 1):
        left = VERSION_ORDER[index]
        right = VERSION_ORDER[index + 1]
        pair_changed: set[str] = set()
        for from_version, to_version in ((left, right), (right, left)):
            mapping = load_direct_mapping(from_version, to_version)
            observed[from_version].update(mapping.keys())
            for source_code, targets in mapping.items():
                for target in targets:
                    target_code = target["code"]
                    observed[to_version].add(target_code)
                    if source_code != target_code:
                        pair_changed.add(source_code)
                        pair_changed.add(target_code)
        changed[(left, right)] = pair_changed

    return observed, changed


@lru_cache(maxsize=4096)
def infer_source_versions(code: str) -> tuple[str, ...]:
    """Return candidate HS versions, including sparse-table identity carryover."""
    observed, changed = build_version_membership_index()
    candidates = {version for version, codes in observed.items() if code in codes}

    # If a code is present on one side of an adjacent boundary and that code
    # has no explicit change on the boundary, the sparse table supports an
    # unchanged carryover to the other side.
    expanded = True
    while expanded:
        expanded = False
        for index in range(len(VERSION_ORDER) - 1):
            left = VERSION_ORDER[index]
            right = VERSION_ORDER[index + 1]
            if code in changed[(left, right)]:
                continue
            if left in candidates or right in candidates:
                if left not in candidates or right not in candidates:
                    candidates.update({left, right})
                    expanded = True

    return tuple(version for version in VERSION_ORDER if version in candidates)


def map_one_step(codes: list[str], from_version: str, to_version: str) -> tuple[list[str], list[dict[str, Any]], list[str]]:
    direct = load_direct_mapping(from_version, to_version)
    next_codes: list[str] = []
    edges: list[dict[str, Any]] = []
    implicit_codes: list[str] = []
    for code in codes:
        targets = direct.get(code)
        if targets:
            for target in targets:
                next_codes.append(target["code"])
                edges.append(
                    {
                        "from": code,
                        "to": target["code"],
                        "weight": target["weight"],
                        "relation": target["relation"],
                    }
                )
        else:
            next_codes.append(code)
            implicit_codes.append(code)
            edges.append(
                {
                    "from": code,
                    "to": code,
                    "weight": 1.0,
                    "relation": "implicit_identity",
                }
            )
    return sorted(set(next_codes)), edges, sorted(set(implicit_codes))


def build_conversion(code: str, source_version: str | None, target_year: int) -> dict[str, Any]:
    target_version = year_to_version(target_year)
    normalized = normalize_code(code)
    requested_source_version = version_from_value(source_version)
    source_candidates = infer_source_versions(normalized)
    source_version_inferred = requested_source_version is None

    if requested_source_version is None:
        if not source_candidates:
            raise ConversionError(
                f"无法从当前转换关系自动识别 {normalized} 的来源 HS 版本，请手动选择来源版本。"
            )
        source_version = source_candidates[0]
    else:
        source_version = requested_source_version
        if source_candidates and source_version not in source_candidates:
            candidate_labels = "、".join(VERSION_LABELS[version] for version in source_candidates)
            raise ConversionError(
                f"{normalized} 不属于 {VERSION_LABELS[source_version]} 的识别范围。"
                f"当前转换关系显示它属于：{candidate_labels}。"
            )

    source_index = VERSION_ORDER.index(source_version)

    codes_by_version: dict[str, list[str]] = {version: [] for version in VERSION_ORDER}
    codes_by_version[source_version] = [normalized]
    edges: list[dict[str, Any]] = []
    implicit_identity: list[dict[str, str]] = []

    current_codes = [normalized]
    for index in range(source_index - 1, -1, -1):
        from_version = VERSION_ORDER[index + 1]
        to_version = VERSION_ORDER[index]
        current_codes, step_edges, implicit = map_one_step(current_codes, from_version, to_version)
        codes_by_version[to_version] = current_codes
        edges.extend({"from_version": from_version, "to_version": to_version, **edge} for edge in step_edges)
        implicit_identity.extend({"version": from_version, "code": code} for code in implicit)

    current_codes = [normalized]
    for index in range(source_index, len(VERSION_ORDER) - 1):
        from_version = VERSION_ORDER[index]
        to_version = VERSION_ORDER[index + 1]
        current_codes, step_edges, implicit = map_one_step(current_codes, from_version, to_version)
        codes_by_version[to_version] = current_codes
        edges.extend({"from_version": from_version, "to_version": to_version, **edge} for edge in step_edges)
        implicit_identity.extend({"version": from_version, "code": code} for code in implicit)

    target_codes = codes_by_version[target_version]
    return {
        "input_code": normalized,
        "source_version": source_version,
        "source_label": VERSION_LABELS[source_version],
        "source_version_inferred": source_version_inferred,
        "source_candidates": list(source_candidates),
        "target_year": target_year,
        "target_version": target_version,
        "target_label": VERSION_LABELS[target_version],
        "versions": [
            {
                "key": version,
                "label": VERSION_LABELS[version],
                "release_year": VERSION_YEARS[version],
                "codes": codes_by_version[version],
            }
            for version in VERSION_ORDER
        ],
        "target_codes": target_codes,
        "edges": edges,
        "implicit_identity": implicit_identity,
        "sources": [
            {"label": "UN Statistics classification", "url": SOURCE_LINKS["un_statistics"]},
            {"label": "WTO HS Tracker", "url": SOURCE_LINKS["wto_hs_tracker"]},
            {"label": "Harvard conversion weights", "path": str(WEIGHTS_DIR)},
        ],
    }


def build_multi_year_conversion(code: str, source_version: str | None, target_years: list[int]) -> dict[str, Any]:
    results = [
        build_conversion(code=code, source_version=source_version, target_year=target_year)
        for target_year in target_years
    ]
    result = results[0]
    result["target_years"] = target_years
    result["target_results"] = [
        {
            "target_year": item["target_year"],
            "target_version": item["target_version"],
            "target_label": item["target_label"],
            "target_codes": item["target_codes"],
        }
        for item in results
    ]
    return result


@lru_cache(maxsize=8)
def load_official_english_descriptions(version: str) -> dict[str, str]:
    """Load six-digit English descriptions from the official UNSD workbook."""
    if version not in VERSION_ORDER:
        raise ConversionError(f"不支持的 HS 版本：{version}")
    if not OFFICIAL_DESCRIPTION_WORKBOOK.exists():
        raise ConversionError(f"缺少官方 HS 描述文件：{OFFICIAL_DESCRIPTION_WORKBOOK.name}")
    try:
        frame = pd.read_excel(
            OFFICIAL_DESCRIPTION_WORKBOOK,
            sheet_name=VERSION_LABELS[version],
            dtype=str,
            engine="openpyxl",
            usecols=["Code", "Description", "Level", "IsBasicLevel"],
        )
    except Exception as exc:  # pragma: no cover - surfaced by the export endpoint
        raise ConversionError(f"读取官方 HS 描述文件失败：{exc}") from exc

    descriptions: dict[str, str] = {}
    for _, row in frame.iterrows():
        raw_code = row.get("Code")
        if pd.isna(raw_code):
            continue
        raw_code = str(raw_code).strip()
        if not re.fullmatch(r"\d{6}", raw_code):
            continue
        try:
            code = normalize_code(str(raw_code))
        except ConversionError:
            continue
        description = row.get("Description")
        if not pd.isna(description):
            descriptions[code] = str(description).strip()
    return descriptions


def descriptions_for_code(code: str, version: str) -> tuple[str, str]:
    """Return the exact target-version English description and an empty Chinese field.

    UNSD's official all-version workbook provides English descriptions. It does
    not provide a universal Chinese HS description table. A Chinese tariff
    table from one country/version cannot be substituted for another HS
    version, so the export keeps the Chinese column empty until a matching
    official source is available for that exact version.
    """
    english = load_official_english_descriptions(version).get(normalize_code(code), "")
    return english, ""


def export_rows(results: list[dict[str, Any]], target_year: int) -> tuple[str, list[list[str]]]:
    target_version = year_to_version(target_year)
    version_name = f"HS{VERSION_YEARS[target_version]}"
    codes: list[str] = []
    for result in results:
        for mapping in result.get("target_results", []):
            if mapping["target_year"] != target_year:
                continue
            for code in mapping.get("target_codes", []):
                if code not in codes:
                    codes.append(code)

    rows = []
    for code in codes:
        english, chinese = descriptions_for_code(code, target_version)
        rows.append([code, english, chinese])
    return version_name, rows


def create_export_workbook(version_name: str, rows: list[list[str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "HS映射结果"
    sheet.append([f"{version_name} code", "含铜产品英文描述", "含铜产品中文描述"])
    for row in rows:
        sheet.append(row)

    header_fill = PatternFill(fill_type="solid", fgColor="E8F1FC")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 58
    sheet.column_dimensions["C"].width = 58

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def parse_conversion_payload(payload: dict[str, Any]) -> tuple[list[str], list[int], list[dict[str, Any]]]:
    raw_codes = payload.get("codes", payload.get("code"))
    codes = parse_hs_codes(raw_codes)
    raw_years = payload.get("target_years", payload.get("target_year"))
    target_years = parse_target_years(raw_years)
    source_version = str(payload.get("source_version", "AUTO"))
    results = [
        build_multi_year_conversion(code, source_version, target_years)
        for code in codes
    ]
    return codes, target_years, results


app = Flask(__name__)


@app.get("/")
def index():
    return render_template(
        "index.html",
        weights_dir=str(WEIGHTS_DIR),
    )


@app.get("/api/health")
def health():
    return jsonify({"ok": WEIGHTS_DIR.is_dir(), "weights_dir": str(WEIGHTS_DIR)})


@app.get("/favicon.ico")
def favicon():
    return "", 204


@app.post("/api/upload-hscodes")
def upload_hscodes():
    try:
        file_storage = request.files.get("file")
        if file_storage is None or not file_storage.filename:
            raise ConversionError("请选择要上传的 CSV 或 XLSX 文件。")
        codes = read_uploaded_hs_codes(file_storage)
        return jsonify({"ok": True, "codes": codes})
    except (ConversionError, ValueError) as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400
    except Exception as exc:  # pragma: no cover - defensive upload boundary
        app.logger.exception("HS code upload failed")
        return jsonify({"ok": False, "error": f"文件处理失败：{exc}"}), 400


@app.post("/api/convert")
def convert():
    payload = request.get_json(silent=True) or {}
    try:
        codes, target_years, results = parse_conversion_payload(payload)
        return jsonify({
            "ok": True,
            "result": {
                "input_codes": codes,
                "target_years": target_years,
                "results": results,
            },
        })
    except (ConversionError, ValueError) as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400
    except Exception as exc:  # pragma: no cover - defensive API boundary
        app.logger.exception("HS conversion failed")
        return jsonify({"ok": False, "error": f"转换失败：{exc}"}), 500


@app.post("/api/export-excel")
def export_excel():
    payload = request.get_json(silent=True) or {}
    try:
        _, target_years, results = parse_conversion_payload(payload)
        files: list[tuple[str, bytes]] = []
        for target_year in target_years:
            version_name, rows = export_rows(results, target_year)
            filename = f"{target_year}_{version_name}.xlsx"
            files.append((filename, create_export_workbook(version_name, rows)))

        if len(files) == 1:
            filename, content = files[0]
            return send_file(
                io.BytesIO(content),
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                as_attachment=True,
                download_name=filename,
            )

        archive = io.BytesIO()
        with ZipFile(archive, "w", compression=ZIP_DEFLATED) as zip_file:
            for filename, content in files:
                zip_file.writestr(filename, content)
        archive.seek(0)
        return send_file(
            archive,
            mimetype="application/zip",
            as_attachment=True,
            download_name="HS_Code_映射结果.zip",
        )
    except (ConversionError, ValueError) as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400
    except Exception as exc:  # pragma: no cover - defensive export boundary
        app.logger.exception("HS Excel export failed")
        return jsonify({"ok": False, "error": f"导出失败：{exc}"}), 500


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=int(os.getenv("PORT", "8010")), debug=True)
