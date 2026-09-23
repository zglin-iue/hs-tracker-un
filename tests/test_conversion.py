import unittest
from io import BytesIO
from zipfile import ZipFile

import pandas as pd
from openpyxl import load_workbook

from app import (
    ConversionError,
    app as flask_app,
    build_conversion,
    parse_hs_codes,
    parse_target_years,
    year_to_version,
)


class ConversionTests(unittest.TestCase):
    def test_target_year_versions(self):
        self.assertEqual(year_to_version(2005), "H2")
        self.assertEqual(year_to_version(2010), "H3")
        self.assertEqual(year_to_version(2015), "H4")
        self.assertEqual(year_to_version(2020), "H5")

    def test_parse_multiple_target_years(self):
        self.assertEqual(parse_target_years("2005 2010 2015 2020"), [2005, 2010, 2015, 2020])
        self.assertEqual(parse_target_years("2005, 2010 2005"), [2005, 2010])

    def test_parse_multiple_hs_codes(self):
        self.assertEqual(parse_hs_codes("010110 390760"), ["010110", "390760"])

    def test_upload_csv_extracts_hs_codes(self):
        with flask_app.test_client() as client:
            response = client.post(
                "/api/upload-hscodes",
                data={"file": (BytesIO(b"HSCode\n010110\n390760\n"), "codes.csv")},
                content_type="multipart/form-data",
            )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["codes"], ["010110", "390760"])

    def test_upload_xlsx_extracts_hs_codes(self):
        workbook = BytesIO()
        pd.DataFrame({"HSCode": ["010110", "390760"]}).to_excel(workbook, index=False)
        workbook.seek(0)
        with flask_app.test_client() as client:
            response = client.post(
                "/api/upload-hscodes",
                data={"file": (workbook, "codes.xlsx")},
                content_type="multipart/form-data",
            )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["codes"], ["010110", "390760"])

    def test_upload_xlsx_ignores_sheets_without_hscode(self):
        workbook = BytesIO()
        with pd.ExcelWriter(workbook, engine="openpyxl") as writer:
            pd.DataFrame({"HSCode": ["260300", "740319"]}).to_excel(
                writer, sheet_name="HS22 code list", index=False
            )
            pd.DataFrame({"Country": ["China"], "Region": ["Asia"]}).to_excel(
                writer, sheet_name="Region", index=False
            )
        workbook.seek(0)
        with flask_app.test_client() as client:
            response = client.post(
                "/api/upload-hscodes",
                data={"file": (workbook, "codes.xlsx")},
                content_type="multipart/form-data",
            )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["codes"], ["260300", "740319"])

    def test_upload_requires_exact_hscode_column(self):
        with flask_app.test_client() as client:
            response = client.post(
                "/api/upload-hscodes",
                data={"file": (BytesIO(b"HS Code\n010110\n"), "codes.csv")},
                content_type="multipart/form-data",
            )
        self.assertEqual(response.status_code, 400)
        self.assertIn("HSCode", response.get_json()["error"])

    def test_export_excel_uses_version_specific_official_english_description(self):
        with flask_app.test_client() as client:
            response = client.post(
                "/api/export-excel",
                json={"codes": "260300", "target_years": "2020"},
            )
        self.assertEqual(response.status_code, 200)
        self.assertIn("2020_HS2017.xlsx", response.headers["Content-Disposition"])
        workbook = load_workbook(BytesIO(response.data), read_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        self.assertEqual(rows[0], ("HS2017 code", "含铜产品英文描述", "含铜产品中文描述"))
        self.assertIn(("260300", "Copper ores and concentrates", None), rows)

    def test_multi_year_export_is_zip_with_year_version_files(self):
        with flask_app.test_client() as client:
            response = client.post(
                "/api/export-excel",
                json={"codes": "260300", "target_years": "2005 2020"},
            )
        self.assertEqual(response.status_code, 200)
        self.assertIn("application/zip", response.headers["Content-Type"])
        with ZipFile(BytesIO(response.data)) as archive:
            self.assertEqual(set(archive.namelist()), {"2005_HS2002.xlsx", "2020_HS2017.xlsx"})

    def test_public_hs_label_is_accepted(self):
        result = build_conversion("010110", "HS02", 2005)
        self.assertEqual(result["source_version"], "H2")

    def test_auto_infers_hs12_for_010121(self):
        result = build_conversion("010121", "AUTO", 2020)
        by_version = {item["key"]: item["codes"] for item in result["versions"]}
        self.assertEqual(result["source_version"], "H4")
        self.assertEqual(result["source_candidates"], ["H4", "H5", "H6"])
        self.assertEqual(by_version["H0"], ["010111", "010120"])
        self.assertEqual(by_version["H2"], ["010110"])
        self.assertEqual(by_version["H4"], ["010121"])
        self.assertEqual(by_version["H6"], ["010121"])

    def test_rejects_incompatible_explicit_version(self):
        with self.assertRaisesRegex(ConversionError, "不属于 HS92"):
            build_conversion("010121", "HS92", 2020)

    def test_hs02_example_keeps_all_branches(self):
        result = build_conversion("010110", "H2", 2005)
        by_version = {item["key"]: item["codes"] for item in result["versions"]}
        self.assertEqual(by_version["H0"], ["010111", "010120"])
        self.assertEqual(by_version["H1"], ["010111", "010120"])
        self.assertEqual(by_version["H2"], ["010110"])
        self.assertEqual(by_version["H4"], ["010121", "010130"])
        self.assertEqual(result["target_codes"], ["010110"])


if __name__ == "__main__":
    unittest.main()
